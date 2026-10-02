import os
import csv

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from striprtf.striprtf import rtf_to_text


def load_document(file_path):

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return PyPDFLoader(file_path).load()

    elif extension == ".docx":
        doc = DocxDocument(file_path)

        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        )

        return [Document(page_content=text)]

    elif extension in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
            text = file.read()

        return [Document(page_content=text)]

    elif extension in [".xlsx", ".xls"]:

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True
        )

        documents = []

        for sheet in workbook.worksheets:

            rows = []

            for row in sheet.iter_rows(values_only=True):

                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:
                    rows.append(" | ".join(values))

            text = "\n".join(rows)

            if text:
                documents.append(
                    Document(
                        page_content=f"Sheet: {sheet.title}\n{text}"
                    )
                )

        return documents

    elif extension == ".csv":

        rows = []

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            reader = csv.reader(file)

            for row in reader:
                rows.append(" | ".join(row))

        return [
            Document(
                page_content="\n".join(rows)
            )
        ]

    elif extension == ".pptx":

        presentation = Presentation(file_path)

        documents = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            texts = []

            for shape in slide.shapes:

                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)

            if texts:

                documents.append(
                    Document(
                        page_content=
                        f"Slide {slide_number}\n" +
                        "\n".join(texts)
                    )
                )

        return documents

    elif extension == ".rtf":

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            text = rtf_to_text(file.read())

        return [Document(page_content=text)]

    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )


def process_pdf(file_path):

    documents = load_document(file_path)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200
    )

    chunks = splitter.split_documents(documents)

    return chunks